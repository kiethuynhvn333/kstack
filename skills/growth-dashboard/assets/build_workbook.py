#!/usr/bin/env python3
"""Build the spreadsheet version of the growth dashboard (Excel / Google Sheets).

Usage:  python3 build_workbook.py dashboard-data.json growth-dashboard.xlsx

Reads the same data file as assets/dashboard-template.html (format at the top of that file) and
writes a workbook laid out as in references/spreadsheet-spec.md: Dashboard, Detail, Notes, Data,
Setup, Calc. Everything on Dashboard and Detail is a formula over the Data tab, driven by the
Segment and Period dropdowns, so a refreshed Data tab updates the whole workbook.
Needs openpyxl (pip install openpyxl).
"""
import calendar
import datetime
import json
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

# Tokens from references/ui-spec.md (light theme).
SERIES = ["2F6FD6", "D9822B", "7A5BD6", "1F9E89", "C94F7C", "8A9A2B", "4AA3C9", "A0522D"]
MUTED_SERIES, FG, MUTED, LINE, OK, BAD, WARN, BG = "BDBCB5", "1C1C1A", "6A6964", "E3E2DD", "1D7A4C", "B42318", "9A6700", "F6F6F3"
SERIF, SANS = "Georgia", "Arial"
TOTAL = "total"
hair = Side(style="thin", color=LINE)
BOX = Border(left=hair, right=hair, top=hair, bottom=hair)


def q(text):
    """A string literal inside a formula."""
    return '"' + str(text).replace('"', '""') + '"'


def number_format(fmt, cur):
    c = cur.replace('"', "")
    return {
        "int": "#,##0",
        "percent": "0.0%",
        "money": f'[>=1000000000]"{c}"#,##0.00,,,"B";[>=1000000]"{c}"#,##0.00,,"M";"{c}"#,##0',
        "money_full": f'[<100]"{c}"#,##0.00;"{c}"#,##0',
    }.get(fmt, "#,##0")


def period_meta(p, grain, days):
    if grain == "week":
        d = datetime.date.fromisoformat(p)
        full = 7
        name = f"Week of {d.day} {d.strftime('%b %Y')}"
        short = f"{d.day} {d.strftime('%b')}"
        return (name if days >= full else f"{name} ({days} days)"), short
    y, m = map(int, p.split("-"))
    full = calendar.monthrange(y, m)[1]
    d = datetime.date(y, m, 1)
    name = d.strftime("%b %Y")
    return (name if days >= full else f"{name} (1–{days})"), d.strftime("%b %y"), full


def main(src, out):
    D = json.load(open(src, encoding="utf-8"))
    cur = D.get("currency", "$")
    grain = D.get("grain", "month")
    segments, sources, metrics = D["segments"], D["sources"], D["metrics"]
    M = {m["id"]: m for m in metrics}
    periods = sorted({r["period"] for r in D["rows"]})
    pdays = {p: max((r.get("days") or 0) for r in D["rows"] if r["period"] == p and r["source"] == TOTAL) for p in periods}
    labels, shorts, fulls = {}, {}, {}
    for p in periods:
        meta = period_meta(p, grain, pdays[p])
        labels[p], shorts[p] = meta[0], meta[1]
        fulls[p] = meta[2] if len(meta) > 2 else 7
    complete = [p for p in periods if pdays[p] >= fulls[p]]
    default_period = (complete or periods)[-1]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    detail, notes, data, setup, calc = (wb.create_sheet(n) for n in ["Detail", "Notes", "Data", "Setup", "Calc"])
    for ws in (setup, calc):
        ws.sheet_properties.tabColor = "BDBCB5"

    # ---------- Data: one long table ----------
    data.append(["period", "segment", "source", "days", "metric", "value"])
    for r in D["rows"]:
        for k, v in (r.get("values") or {}).items():
            if v is not None:
                data.append([r["period"], r["segment"], r["source"], r.get("days"), k, v])
    for c, w in zip("ABCDEF", [12, 16, 16, 8, 14, 16]):
        data.column_dimensions[c].width = w
    data.freeze_panes = "A2"

    # ---------- Setup: the lists formulas read ----------
    setup["A1"] = "Segments"
    for i, s in enumerate(segments, 2):
        setup.cell(i, 1, s)
    setup["C1"], setup["D1"], setup["E1"] = "Period", "Label", "Days"
    for i, p in enumerate(periods, 2):
        setup.cell(i, 3, p), setup.cell(i, 4, labels[p]), setup.cell(i, 5, pdays[p])
    setup["G1"], setup["H1"], setup["I1"], setup["J1"] = "Source id", "Name", "Group", "Not linked?"
    for i, s in enumerate(sources, 2):
        setup.cell(i, 7, s["id"]), setup.cell(i, 8, s["name"]), setup.cell(i, 9, s.get("group", "")), setup.cell(i, 10, bool(s.get("unlinked")))
    setup["L1"], setup["M1"], setup["N1"], setup["O1"], setup["P1"], setup["Q1"] = "Metric id", "Name", "What it counts", "Format", "Up is", "Compare per day"
    for i, m in enumerate(metrics, 2):
        for j, k in enumerate(["id", "name", "sub", "format", "good", "perDay"], 12):
            setup.cell(i, j, m.get(k, False if k == "perDay" else ""))
    th = (D.get("status") or {}).get("threshold", 0.05)
    setup["S1"], setup["T1"] = "Setting", "Value"
    for i, (k, v) in enumerate([("Growing / declining threshold", th), ("Currency", cur), ("Data to", datetime.date.fromisoformat(D["asOf"]) if D.get("asOf") else "")], 2):
        setup.cell(i, 19, k), setup.cell(i, 20, v)
    for c in "ACDGHILMN":
        setup.column_dimensions[c].width = 18
    np_ = len(periods)
    P_IDS, P_LAB, P_DAYS = f"Setup!$C$2:$C${np_ + 1}", f"Setup!$D$2:$D${np_ + 1}", f"Setup!$E$2:$E${np_ + 1}"

    # ---------- Calc: selection and helper tables ----------
    SEG, PER, PREV, CDAYS, PDAYS = "Calc!$B$1", "Calc!$B$2", "Calc!$B$3", "Calc!$B$4", "Calc!$B$5"
    calc["A1"], calc["B1"] = "Segment", "=Dashboard!$C$4"
    calc["A2"], calc["B2"] = "Period", f"=INDEX({P_IDS},MATCH(Dashboard!$E$4,{P_LAB},0))"
    calc["A3"], calc["B3"] = "Previous period", f'=IFERROR(INDEX({P_IDS},MATCH({PER},{P_IDS},0)-1),"")'
    calc["A4"], calc["B4"] = "Days", f"=VLOOKUP({PER},Setup!$C:$E,3,FALSE)"
    calc["A5"], calc["B5"] = "Days before", f'=IFERROR(VLOOKUP({PREV},Setup!$C:$E,3,FALSE),"")'

    def V(p, s, src_, m):
        crit = f"Data!$A:$A,{p},Data!$B:$B,{s},Data!$C:$C,{src_},Data!$E:$E,{q(m)}"
        return f'IF(COUNTIFS({crit})=0,"",SUMIFS(Data!$F:$F,{crit}))'

    def CHG(cur_, prev_, m):
        f = M[m]
        if f.get("format") == "percent":
            return f'IF(OR({cur_}="",{prev_}=""),"",({cur_}-{prev_})*100)'
        if f.get("perDay"):
            return f'IFERROR(({cur_}/{CDAYS})/({prev_}/{PDAYS})-1,"")'
        return f'IFERROR({cur_}/{prev_}-1,"")'

    # metric block: rows 8.. (A id, B now, C before, D change)
    calc["A7"], calc["B7"], calc["C7"], calc["D7"] = "Metric", "Now", "Before", "Change"
    mrow = {}
    for i, m in enumerate(metrics, 8):
        mrow[m["id"]] = i
        calc.cell(i, 1, m["id"])
        calc.cell(i, 2, "=" + V(PER, SEG, q(TOTAL), m["id"]))
        calc.cell(i, 3, f'=IF({PREV}="","",{V(PREV, SEG, q(TOTAL), m["id"])})')
        calc.cell(i, 4, "=" + CHG(f"B{i}", f"C{i}", m["id"]))

    # status block: per source now / before / change / top-source helper
    st = D.get("status") or {}
    sm, em = st.get("metric", D["contribution"]["metric"]), st.get("efficiencyMetric")
    srow0 = 8 + len(metrics) + 2
    calc.cell(srow0 - 1, 1, "Source id"), calc.cell(srow0 - 1, 2, "Now"), calc.cell(srow0 - 1, 3, "Before"), calc.cell(srow0 - 1, 4, "Change"), calc.cell(srow0 - 1, 5, "For top source")
    for i, s in enumerate(sources):
        r = srow0 + i
        calc.cell(r, 1, s["id"])
        calc.cell(r, 2, "=" + V(PER, SEG, q(s["id"]), sm))
        calc.cell(r, 3, f'=IF({PREV}="","",{V(PREV, SEG, q(s["id"]), sm)})')
        calc.cell(r, 4, "=" + CHG(f"B{r}", f"C{r}", sm))
        calc.cell(r, 5, -1 if s.get("unlinked") else f"=N(B{r})")
    srow1 = srow0 + len(sources) - 1

    # contribution table (periods x sources), columns G..
    c = D["contribution"]
    g0 = 7
    calc.cell(7, g0, "Period")
    for j, s in enumerate(sources, 1):
        calc.cell(7, g0 + j, s["name"])
    tcol = g0 + len(sources) + 1
    calc.cell(7, tcol, "Total")
    calc.cell(7, tcol + 1, f'{c.get("shareGroup", "")} share')
    share_ids = [s["id"] for s in sources if c.get("shareGroup") and s.get("group") == c["shareGroup"]]
    for i, p in enumerate(periods, 8):
        calc.cell(i, g0, shorts[p])
        pq = q(p)
        for j, s in enumerate(sources, 1):
            calc.cell(i, g0 + j, "=" + V(pq, SEG, q(s["id"]), c["metric"]))
        calc.cell(i, tcol, "=" + V(pq, SEG, q(TOTAL), c["metric"]))
        if share_ids:
            parts = "+".join(f"N({col(g0 + 1 + [s['id'] for s in sources].index(sid))}{i})" for sid in share_ids)
            calc.cell(i, tcol + 1, f'=IFERROR(({parts})/{col(tcol)}{i},"")')
    plast = 7 + len(periods)

    # efficiency table (periods x chosen sources), after the contribution table
    e = D.get("efficiency") or {}  # optional: no spend data, no cost chart
    if not (e.get("metric") and e.get("sources")):
        e = {"sources": []}
    e0 = tcol + 3
    names = {s["id"]: s["name"] for s in sources}
    if e["sources"]:
        calc.cell(7, e0, "Period")
    for j, sid in enumerate(e["sources"], 1):
        calc.cell(7, e0 + j, "Total" if sid == TOTAL else names[sid])
    for i, p in enumerate(periods if e["sources"] else [], 8):
        calc.cell(i, e0, shorts[p])
        for j, sid in enumerate(e["sources"], 1):
            calc.cell(i, e0 + j, f'=IF({V(q(p), SEG, q(sid), e["metric"])}="",NA(),{V(q(p), SEG, q(sid), e["metric"])})')

    # funnel table
    fn = D.get("funnel") or {}
    f0 = e0 + len(e["sources"]) + 2
    if fn.get("steps"):
        calc.cell(7, f0, "Step"), calc.cell(7, f0 + 1, "People")
        for i, m in enumerate(fn["steps"], 8):
            calc.cell(i, f0, M[m]["name"])
            calc.cell(i, f0 + 1, "=" + V(PER, SEG, q(TOTAL), m))

    # ---------- Dashboard ----------
    dash.sheet_view.showGridLines = False
    dash.column_dimensions["A"].width = 2
    for cc in "BCDEFGHI":
        dash.column_dimensions[cc].width = 22
    dash["B1"] = D.get("title", "Growth Overview")
    dash["B1"].font = Font(name=SERIF, size=22, color=FG)
    dash.row_dimensions[1].height = 34

    hm = c["metric"]
    hfmt = M[hm].get("format", "int")
    vtext = {"percent": f'TEXT(Calc!$B${mrow[hm]},"0.0%")', "money": f'{q(cur)}&TEXT(Calc!$B${mrow[hm]},"#,##0")',
             "money_full": f'{q(cur)}&TEXT(Calc!$B${mrow[hm]},"#,##0")'}.get(hfmt, f'TEXT(Calc!$B${mrow[hm]},"#,##0")')
    chg_txt = f'IF(Calc!$D${mrow[hm]}="",""," ("&TEXT(Calc!$D${mrow[hm]},"+0.0%;-0.0%")&" vs the period before)")'
    if hfmt == "percent":
        chg_txt = f'IF(Calc!$D${mrow[hm]}="",""," ("&TEXT(Calc!$D${mrow[hm]},"+0.0;-0.0")&" pts vs the period before)")'
    top = (f'IF(MAX(Calc!$E${srow0}:$E${srow1})<=0,"",". The biggest source was "&INDEX(Setup!$H$2:$H${len(sources) + 1},'
           f'MATCH(MAX(Calc!$E${srow0}:$E${srow1}),Calc!$E${srow0}:$E${srow1},0))&" with "&TEXT(IFERROR(MAX(Calc!$E${srow0}:$E${srow1})/Calc!$B${mrow[hm]},0),"0.0%")&" of it")')
    dash["B2"] = (f'="In "&$E$4&", "&IF($C$4="All","the business",$C$4)&" had "&{vtext}&" "&{q(M[hm].get("phrase") or M[hm]["name"])}'
                  f'&{chg_txt}&{top}&"."')
    dash["B2"].font = Font(name=SANS, size=11, color=MUTED)
    if D.get("sample"):
        dash["B3"] = "Sample — not your data"
        dash["B3"].font = Font(name=SANS, size=10, bold=True, color=WARN)

    dash["B4"], dash["D4"] = "Segment", "Period"
    for a in ("B4", "D4"):
        dash[a].font = Font(name=SANS, size=10, color=MUTED)
        dash[a].alignment = Alignment(horizontal="right")
    dash["C4"], dash["E4"] = segments[0], labels[default_period]
    for a in ("C4", "E4"):
        dash[a].font = Font(name=SANS, size=11, bold=True, color=FG)
        dash[a].border = BOX
        dash[a].fill = PatternFill("solid", fgColor="FFFFFF")
    dv1 = DataValidation(type="list", formula1=f"=Setup!$A$2:$A${len(segments) + 1}", allow_blank=False)
    dv2 = DataValidation(type="list", formula1=f"={P_LAB}", allow_blank=False)
    dash.add_data_validation(dv1), dash.add_data_validation(dv2)
    dv1.add("C4"), dv2.add("E4")
    dash["G4"] = '="Data to "&TEXT(Setup!$T$4,"d mmm yyyy")'
    setup["T4"].number_format = "yyyy-mm-dd"
    dash["G4"].font = Font(name=SANS, size=10, color=MUTED)

    # tiles
    row = 6
    for group in D["tileGroups"]:
        dash.cell(row, 2, group["label"].upper()).font = Font(name=SANS, size=9, color=MUTED)
        for j, mid in enumerate(group["metrics"][:4]):
            cc = 2 + j
            m = M[mid]
            r = mrow[mid]
            name = dash.cell(row + 1, cc, f'{m["name"]} · {m["sub"]}' if m.get("sub") else m["name"])
            name.font = Font(name=SANS, size=9, color=MUTED)
            name.alignment = Alignment(wrap_text=True, vertical="top")
            val = dash.cell(row + 2, cc, f"=Calc!$B${r}")
            val.font = Font(name=SERIF, size=20, color=FG)
            val.number_format = number_format(m.get("format"), cur)
            val.alignment = Alignment(horizontal="left")
            ch = dash.cell(row + 3, cc, f"=Calc!$D${r}")
            ch.font = Font(name=SANS, size=9, color=MUTED)
            ch.alignment = Alignment(horizontal="left")
            if m.get("format") == "percent":
                ch.number_format = '+0.0" pts vs before";-0.0" pts vs before";0.0" pts vs before"'
                small = 0.05
            else:
                word = " per day" if m.get("perDay") else ""
                ch.number_format = f'+0.0%"{word} vs before";-0.0%"{word} vs before";0.0%"{word} vs before"'
                small = 0.005
            ref = ch.coordinate
            good = m.get("good", "up")
            if good != "neutral":
                up, down = (OK, BAD) if good == "up" else (BAD, OK)
                dash.conditional_formatting.add(ref, FormulaRule(formula=[f"AND(ISNUMBER({ref}),{ref}>{small})"], font=Font(color=up)))
                dash.conditional_formatting.add(ref, FormulaRule(formula=[f"AND(ISNUMBER({ref}),{ref}<-{small})"], font=Font(color=down)))
            for rr in range(row + 1, row + 4):
                cell = dash.cell(rr, cc)
                cell.border = Border(left=hair, right=hair, top=hair if rr == row + 1 else None, bottom=hair if rr == row + 3 else None)
                cell.fill = PatternFill("solid", fgColor="FFFFFF")
        dash.row_dimensions[row + 1].height = 34
        dash.row_dimensions[row + 2].height = 30
        row += 5

    # charts
    chart_row = row + 1

    def style_axes(ch):
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.legend.position = "b"

    bar = BarChart()
    bar.type, bar.grouping, bar.overlap, bar.gapWidth = "col", "stacked", 100, 60
    bar.title = c["title"]
    bar.add_data(Reference(calc, min_col=g0 + 1, max_col=g0 + len(sources), min_row=7, max_row=plast), titles_from_data=True)
    bar.set_categories(Reference(calc, min_col=g0, min_row=8, max_row=plast))
    for i, s in enumerate(bar.series):
        color = MUTED_SERIES if sources[i].get("unlinked") else SERIES[i % 8]
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    bar.y_axis.number_format = number_format(M[c["metric"]].get("format"), cur)
    bar.width, bar.height = 17, 8.5
    style_axes(bar)
    dash.add_chart(bar, f"B{chart_row}")

    if e["sources"]:
        line = LineChart()
        line.title = e.get("title", "")
        line.add_data(Reference(calc, min_col=e0 + 1, max_col=e0 + len(e["sources"]), min_row=7, max_row=plast), titles_from_data=True)
        line.set_categories(Reference(calc, min_col=e0, min_row=8, max_row=plast))
        for i, s in enumerate(line.series):
            sid = e["sources"][i]
            color = FG if sid == TOTAL else (MUTED_SERIES if next(x for x in sources if x["id"] == sid).get("unlinked") else SERIES[[x["id"] for x in sources].index(sid) % 8])
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 25400
            s.marker.symbol = "circle"
            s.marker.size = 5
            s.marker.graphicalProperties.solidFill = color
            s.marker.graphicalProperties.line.solidFill = color
            s.smooth = False
        line.y_axis.number_format = number_format(M[e["metric"]].get("format"), cur).replace("0.00;", "0;")
        line.display_blanks = "gap"
        line.width, line.height = 17, 8.5
        style_axes(line)
        dash.add_chart(line, f"F{chart_row}")

    next_row = chart_row + 18
    if fn.get("steps"):
        fb = BarChart()
        fb.type, fb.gapWidth = "bar", 40
        fb.title = fn.get("title", "Funnel")
        fb.add_data(Reference(calc, min_col=f0 + 1, min_row=7, max_row=7 + len(fn["steps"])), titles_from_data=True)
        fb.set_categories(Reference(calc, min_col=f0, min_row=8, max_row=7 + len(fn["steps"])))
        fb.series[0].graphicalProperties.solidFill = SERIES[0]
        fb.series[0].graphicalProperties.line.solidFill = SERIES[0]
        fb.x_axis.scaling.orientation = "maxMin"
        fb.y_axis.number_format = "#,##0"
        fb.legend = None
        fb.x_axis.delete, fb.y_axis.delete = False, False
        fb.width, fb.height = 17, 2.2 + 1.4 * len(fn["steps"])
        dash.add_chart(fb, f"B{next_row}")
        next_row += int(5 + 3 * len(fn["steps"]))

    # status of each source
    r0 = next_row
    dash.cell(r0, 2, st.get("title", "Status of each source of growth")).font = Font(name=SANS, size=13, bold=True, color=FG)
    dash.cell(r0 + 1, 2, f"Change vs the period before. Growing or declining = more than {th:.0%} either way.").font = Font(name=SANS, size=9, color=MUTED)
    heads = ["Source", "Group", M[sm]["name"], "Share of total", "Change", "Status"] + ([M[em]["name"]] if em else [])
    for j, h in enumerate(heads):
        cell = dash.cell(r0 + 3, 2 + j, h)
        cell.font = Font(name=SANS, size=9, bold=True, color=MUTED)
        cell.border = Border(bottom=hair)
        cell.alignment = Alignment(horizontal="left" if j < 2 else "right")
    for i, s in enumerate(sources):
        r, cr = r0 + 4 + i, srow0 + i
        vals = [s["name"], s.get("group", ""), f"=Calc!$B${cr}", f'=IFERROR(Calc!$B${cr}/Calc!$B${mrow[sm]},"")', f"=Calc!$D${cr}",
                f'=IF(N(Calc!$B${cr})=0,"No data",IF(N(Calc!$C${cr})=0,"New",IF(N(Calc!$D${cr})>Setup!$T$2,"Growing",IF(N(Calc!$D${cr})<-Setup!$T$2,"Declining","Flat"))))']
        if em:
            vals.append("=" + V(PER, SEG, q(s["id"]), em))
        for j, v in enumerate(vals):
            cell = dash.cell(r, 2 + j, v)
            cell.font = Font(name=SANS, size=10, color=FG)
            cell.border = Border(bottom=hair)
            cell.alignment = Alignment(horizontal="left" if j < 2 else "right")
        dash.cell(r, 4).number_format = number_format(M[sm].get("format"), cur)
        dash.cell(r, 5).number_format = "0.0%"
        dash.cell(r, 6).number_format = '+0.0%;-0.0%;0.0%' if M[sm].get("format") != "percent" else '+0.0" pts";-0.0" pts"'
        if em:
            dash.cell(r, 8).number_format = number_format(M[em].get("format"), cur)
    r1 = r0 + 3 + len(sources)
    dash.conditional_formatting.add(f"E{r0 + 4}:E{r1}", DataBarRule(start_type="num", start_value=0, end_type="max", color=SERIES[0], showValue=True))
    chg_rng = f"F{r0 + 4}:F{r1}"
    dash.conditional_formatting.add(chg_rng, FormulaRule(formula=[f"AND(ISNUMBER(F{r0 + 4}),F{r0 + 4}>Setup!$T$2)"], font=Font(color=OK)))
    dash.conditional_formatting.add(chg_rng, FormulaRule(formula=[f"AND(ISNUMBER(F{r0 + 4}),F{r0 + 4}<-Setup!$T$2)"], font=Font(color=BAD)))
    stat_rng = f"G{r0 + 4}:G{r1}"
    for word, color in [("Growing", OK), ("Declining", BAD), ("New", WARN), ("Flat", MUTED), ("No data", MUTED)]:
        dash.conditional_formatting.add(stat_rng, CellIsRule(operator="equal", formula=[q(word)], font=Font(color=color, bold=True)))

    # ---------- Detail ----------
    detail.sheet_view.showGridLines = False
    detail["A1"] = '="All segments, "&Dashboard!$E$4'
    detail["A1"].font = Font(name=SERIF, size=16, color=FG)
    detail["A2"] = "Every segment and source for the period picked on the Dashboard. Bold rows are totals."
    detail["A2"].font = Font(name=SANS, size=9, color=MUTED)
    tm = (D.get("table") or {}).get("metrics") or [m["id"] for m in metrics]
    for j, h in enumerate(["Segment", "Source"] + [M[m]["name"] for m in tm], 1):
        cell = detail.cell(4, j, h)
        cell.font = Font(name=SANS, size=9, bold=True, color=MUTED)
        cell.border = Border(bottom=hair)
        cell.alignment = Alignment(horizontal="left" if j <= 2 else "right", wrap_text=True)
    r = 5
    for seg in segments:
        for sid in [TOTAL] + [s["id"] for s in sources]:
            is_total = sid == TOTAL
            detail.cell(r, 1, seg), detail.cell(r, 2, "Total" if is_total else names[sid])
            for j, m in enumerate(tm, 3):
                cell = detail.cell(r, j, "=" + V(PER, q(seg), q(sid), m))
                cell.number_format = number_format(M[m].get("format"), cur)
            for j in range(1, len(tm) + 3):
                cell = detail.cell(r, j)
                cell.font = Font(name=SANS, size=10, bold=is_total, color=FG if (is_total or j > 1) else MUTED)
                cell.border = Border(bottom=hair)
            r += 1
    detail.column_dimensions["A"].width = 16
    detail.column_dimensions["B"].width = 16
    for j in range(3, len(tm) + 3):
        detail.column_dimensions[col(j)].width = 15
    detail.freeze_panes = "C5"

    # ---------- Notes ----------
    notes.sheet_view.showGridLines = False
    notes["A1"] = "How to read this"
    notes["A1"].font = Font(name=SERIF, size=16, color=FG)
    for i, n in enumerate(D.get("notes", []), 3):
        cell = notes.cell(i, 1, "• " + n)
        cell.font = Font(name=SANS, size=10, color=MUTED)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes.column_dimensions["A"].width = 110

    for ws in (dash, detail, notes):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0
    wb.save(out)
    print(f"Wrote {out}: {len(periods)} periods, {len(segments)} segments, {len(sources)} sources, {data.max_row - 1} data rows")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
